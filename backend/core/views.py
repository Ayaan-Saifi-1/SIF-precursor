import csv,io,json,os,zipfile
from datetime import timedelta
from django.conf import settings
from django.db import transaction,IntegrityError
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework.decorators import api_view
from rest_framework.response import Response
from rest_framework.exceptions import ValidationError
from .models import Report,Analysis,ReviewItem,ReviewDecision,AuditEvent
from .serializers import IntakeSerializer,DecisionSerializer
from .services import run_analysis,report_data,public_analysis
from ml.density import density
from ml.pattern_engine import discover
from ml.encoder import MODEL_DIR
from ml.sif_classifier import load_classifier
from data.demo_seed import SAMPLES

def window_rows(request, apply_filters=True):
    try: days=int(request.query_params.get("days","60"))
    except ValueError: raise ValidationError({"days":"Must be 7 to 365."})
    if not 7<=days<=365: raise ValidationError({"days":"Must be 7 to 365."})
    reports=Report.objects.filter(event_timestamp__gte=timezone.now()-timedelta(days=days)).prefetch_related("analyses")
    if request.query_params.get("demo") == "false":
        reports=reports.filter(is_synthetic=False)
    site=request.query_params.get("site") if apply_filters else None
    if site: reports=reports.filter(site=site)
    rows=[r.analyses.all()[0].payload for r in reports if r.analyses.all()]
    if apply_filters:
        for key in ("activity", "report_type"):
            value=request.query_params.get(key)
            if value: rows=[row for row in rows if row.get(key)==value]
    return rows

def enriched(analysis):
    result=public_analysis(analysis)
    family_ids=result.get("family_ids",[])
    related=[]
    for r in Report.objects.exclude(pk=analysis.report_id).prefetch_related("analyses"):
        if r.analyses.all():
            row=r.analyses.all()[0].payload
            if not row.get("simulated") and set(family_ids)&set(row.get("family_ids",[])):
                related.append(row)
    result.update(pattern_id=family_ids[0] if family_ids and not result.get("simulated") else None,
                  similar_report_count=len(related),sites_affected=sorted(set(r["site"] for r in related)))
    return result

@api_view(["GET"])
def health(request):
    # A manifest on disk does not prove that the artifact is loadable or valid.
    try:
        classifier, manifest = load_classifier()
        status = "READY" if classifier is not None else "NOT_TRAINED"
    except Exception:
        manifest, status = None, "UNAVAILABLE"
    return Response(dict(service="ASCENSION",mode="LOCAL_DEMO" if settings.DEMO_MODE else "AUTHENTICATED",
        encoder_status="INSTALLED" if (MODEL_DIR/"model.safetensors").exists() else "NOT_INSTALLED",
        classifier_status=status,model_version=(manifest or {}).get("model_version"),
        development_only=(manifest or {}).get("development_only",False),
        data_notice="Local model results require HSE review before operational use."))

@api_view(["GET"])
def samples(request): return Response(SAMPLES)

def create_report(data):
    serializer=IntakeSerializer(data=data); serializer.is_valid(raise_exception=True)
    try:
        with transaction.atomic():
            report=serializer.save()
            AuditEvent.objects.create(report=report,event_type="REPORT_INGESTED",payload=report.source_snapshot)
    except IntegrityError:
        raise ValidationError({"report_id":"This report ID already exists."})
    return report

@api_view(["GET","POST"])
def reports(request):
    if request.method=="POST":
        report=create_report(request.data)
        return Response(report_data(report),status=201)
    return Response([dict(**report_data(r),created_at=r.created_at.isoformat()) for r in Report.objects.all()[:500]])

@api_view(["GET"])
def report_detail(request,report_id):
    return Response(report_data(get_object_or_404(Report,report_id=report_id)))

@api_view(["POST"])
def analyze_report(request):
    # Accept a new canonical report or an existing ID for an appended assessment.
    if set(request.data)=={"report_id"}:
        report=get_object_or_404(Report,report_id=request.data["report_id"])
    else: report=create_report(request.data)
    a=run_analysis(report)
    return Response(enriched(a),status=201)

@api_view(["GET"])
def analysis_detail(request,report_id):
    report=get_object_or_404(Report,report_id=report_id)
    analysis=report.analyses.first()
    if not analysis: return Response({"detail":"Report retained; analysis pending."},status=202)
    return Response(enriched(analysis))

@api_view(["POST"])
def import_reports(request):
    upload=request.FILES.get("file")
    if not upload: raise ValidationError({"file":"Upload a CSV or XLSX file."})
    if upload.size>5*1024*1024: raise ValidationError({"file":"Maximum upload size is 5 MB."})
    try:
        if upload.name.lower().endswith(".csv"):
            rows=list(csv.DictReader(io.StringIO(upload.read().decode("utf-8-sig"))))
        elif upload.name.lower().endswith(".xlsx"):
            with zipfile.ZipFile(upload) as z:
                if sum(f.file_size for f in z.infolist())>25*1024*1024:
                    raise ValueError("Expanded workbook exceeds 25 MB.")
            upload.seek(0)
            from openpyxl import load_workbook
            wb=load_workbook(upload,read_only=True,data_only=True)
            try:
                it=iter(wb.active.iter_rows(values_only=True)); headers=next(it)
                rows=[]
                for values in it:
                    if len(rows)>=100: raise ValueError("Maximum 100 reports per import.")
                    if any(v is not None for v in values):
                        rows.append({str(h):v.isoformat() if hasattr(v,"isoformat") else v for h,v in zip(headers,values) if h is not None and v is not None})
            finally: wb.close()
        else: raise ValueError("Only .csv and .xlsx files are supported.")
    except (ValueError,UnicodeDecodeError,zipfile.BadZipFile,StopIteration) as exc:
        raise ValidationError({"file":str(exc)})
    if not rows or len(rows)>100: raise ValidationError({"file":"Import between 1 and 100 reports."})
    seen=set(); valid=[]; errors=[]
    for i,row in enumerate(rows,2):
        s=IntakeSerializer(data=row)
        if not s.is_valid(): errors.append(dict(row=i,errors=s.errors))
        elif row["report_id"] in seen: errors.append(dict(row=i,errors={"report_id":"Duplicate within this file."}))
        else: valid.append(s)
        seen.add(row.get("report_id"))
    if errors: return Response(dict(errors=errors,imported=0),status=400)
    try:
        with transaction.atomic():
            saved=[s.save() for s in valid]
            for r in saved: AuditEvent.objects.create(report=r,event_type="REPORT_IMPORTED",payload=r.source_snapshot)
    except IntegrityError: raise ValidationError({"file":"A report ID was created concurrently. No records imported."})
    results=[public_analysis(run_analysis(r)) for r in saved]
    return Response(dict(imported=len(saved),analyses=results),status=201)

@api_view(["GET"])
def summary(request):
    rows=window_rows(request)
    ids=[r["report_id"] for r in rows]
    return Response(dict(reports_analysed=len(rows),sif_potential=sum(r["sif_label"]=="SIF_POTENTIAL" for r in rows),
        critical_barrier_failures=sum(any(b["state"] in ["BYPASSED","FAILED","ABSENT"] for b in r["barriers"]) for r in rows),
        review_required=ReviewItem.objects.filter(report__report_id__in=ids).exclude(status="REVIEWED").count(),
        unavailable=sum(r["status"]=="ANALYSIS_UNAVAILABLE" for r in rows),
        synthetic_count=sum(r.get("is_synthetic",False) for r in rows),
        sites=sorted(set(r["site"] for r in rows)),
        recent_critical=[{k:v for k,v in r.items() if k!="embedding"} for r in rows if r["priority"] in ["CRITICAL","HIGH"]][:5]))
@api_view(["GET"])
def site_density(request): return Response(density(window_rows(request),"site"))
@api_view(["GET"])
def activity_density(request): return Response(density(window_rows(request),"activity"))
@api_view(["GET"])
def dashboard_charts(request):
    from .dashboard_data import chart_data
    rows = window_rows(request)
    result=chart_data(rows, int(request.query_params.get("days", "60")))
    available=window_rows(request, apply_filters=False)
    result["filter_options"]={key: sorted({r.get(key) or "Not identified" for r in available})
                              for key in ("site", "activity", "report_type")}
    return Response(result)
@api_view(["GET"])
def patterns(request): return Response(discover(window_rows(request)))
@api_view(["GET"])
def pattern_detail(request,pattern_id):
    p=next((p for p in discover(window_rows(request)) if p["pattern_id"]==pattern_id),None)
    return Response(p if p else {"detail":"Pattern not found"},status=200 if p else 404)
@api_view(["GET"])
def pattern_reports(request,pattern_id):
    rows=window_rows(request)
    return Response([{k:v for k,v in r.items() if k!="embedding"} for r in rows if pattern_id in r.get("family_ids",[]) and not r.get("simulated")])

def review_json(item):
    return dict(id=item.id,report_id=item.report.report_id,site=item.report.site,description=item.report.description,
        priority=item.priority,sif_label=item.analysis.sif_label,barrier_states=[b["state"] for b in item.analysis.payload["barriers"]],
        review_reason=item.review_reason,status=item.status,version=item.version,created_at=item.created_at.isoformat(),
        decisions=[dict(decision=d.decision,reviewer=d.reviewer,identity_verified=d.identity_verified,
        override_reason=d.override_reason,timestamp=d.created_at.isoformat()) for d in item.decisions.all()])
@api_view(["GET"])
def reviews(request):
    items=ReviewItem.objects.select_related("report","analysis").prefetch_related("decisions")
    status=request.query_params.get("status")
    if status and status!="ALL": items=items.filter(status=status)
    order={"CRITICAL":0,"HIGH":1,"MEDIUM":2,"LOW":3}
    return Response([review_json(i) for i in sorted(items,key=lambda i:(order[i.priority],i.created_at))])
@api_view(["GET"])
def review_detail(request,review_id): return Response(review_json(get_object_or_404(ReviewItem,pk=review_id)))
@api_view(["POST"])
def review_decision(request,review_id):
    s=DecisionSerializer(data=request.data); s.is_valid(raise_exception=True); data=s.validated_data
    with transaction.atomic():
        item=get_object_or_404(ReviewItem.objects.select_for_update(),pk=review_id)
        if item.version!=data["version"]: return Response({"detail":"This review changed. Refresh before submitting."},status=409)
        verified=request.user.is_authenticated
        reviewer=request.user.get_username() if verified else data["reviewer"]
        d=ReviewDecision.objects.create(review=item,decision=data["decision"],reviewer=reviewer,identity_verified=verified,
            override_reason=data["override_reason"],ai_recommendation=public_analysis(item.analysis))
        item.status={"CONFIRM":"REVIEWED","DISAGREE":"REVIEWED","REQUEST_MORE_INFORMATION":"AWAITING_INFORMATION","ESCALATE":"ESCALATED"}[d.decision]
        item.version+=1; item.save(update_fields=["status","version"])
        AuditEvent.objects.create(report=item.report,event_type="HSE_DECISION",payload=dict(decision_id=d.id,
            decision=d.decision,reviewer=reviewer,identity_verified=verified,override_reason=d.override_reason,
            analysis_id=item.analysis_id,review_status=item.status))
    return Response(review_json(item))
